from __future__ import annotations

import io
import json
import re
import sqlite3
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse

try:
    from pypdf import PdfReader
except ImportError:  # pragma: no cover
    PdfReader = None

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from fpdf import FPDF
except ImportError:  # pragma: no cover
    FPDF = None

app = FastAPI()
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
AUDIT_DB = DATA_DIR / "auditquest.db"
AUDIT_STORE = DATA_DIR / "audits.json"


def _iso_timestamp() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _extract_marks(text: str) -> int:
    matches = re.findall(r"\((\d+)\s*marks?\)|\[(\d+)\]|\b(\d+)\s*marks?\b", text, flags=re.I)
    flat = [int(v or w or x) for group in matches for v, w, x in [group] if (v or w or x)]
    return max(flat) if flat else 10


def extract_text(filename: str, content: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext in {".txt", ".md"}:
        return content.decode("utf-8", errors="replace")
    if ext == ".pdf":
        if not PdfReader:
            raise HTTPException(500, "pypdf not installed")
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if ext == ".docx":
        if not Document:
            raise HTTPException(500, "python-docx not installed")
        return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs)
    raise HTTPException(415, "Unsupported format")


def classify_bloom(t: str):
    text = t.lower()
    weights = {
        "Remember": 0,
        "Understand": 0,
        "Apply": 0,
        "Analyze": 0,
        "Evaluate": 0,
        "Create": 0,
    }
    rules = [
        ("Create", ["design", "create", "develop", "plan", "propose", "construct"]),
        ("Evaluate", ["justify", "assess", "critique", "evaluate", "recommend", "judge"]),
        ("Analyze", ["compare", "analyze", "differentiate", "infer", "contrast", "examine"]),
        ("Apply", ["solve", "calculate", "apply", "use", "demonstrate", "implement"]),
        ("Understand", ["explain", "describe", "interpret", "define", "summarize", "elaborate"]),
        ("Remember", ["list", "state", "identify", "remember", "define", "name"]),
    ]
    for level, words in rules:
        weights[level] = sum(text.count(word) for word in words)
    if max(weights.values()) == 0:
        return "Understand"
    return max(weights, key=weights.get)


def classify_diff(t: str, bloom_level: str):
    text = t.lower()
    if bloom_level in {"Create", "Evaluate"}:
        return "Hard"
    if len(t.split()) > 20 or "compare" in text or "differentiate" in text:
        return "Medium"
    return "Easy"


def parse_questions(text: str):
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    pattern = re.compile(r"^(?:Q|Question)?\s*(\d+)[.)]\s*(.+)$", re.I)
    qs: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None

    for line in lines:
        match = pattern.match(line)
        if match:
            if current:
                qs.append(current)
            current = {"number": match.group(1), "text": match.group(2).strip()}
        elif current:
            current["text"] += " " + line

    if current:
        qs.append(current)

    if not qs:
        for index, line in enumerate(lines, start=1):
            if "?" in line or len(line) > 20:
                qs.append({"number": index, "text": line})

    if not qs:
        raise ValueError("No questions detected")

    for q in qs:
        q["text"] = q["text"].strip()
        q["marks"] = _extract_marks(q["text"])
        q["bloom"] = classify_bloom(q["text"])
        q["difficulty"] = classify_diff(q["text"], q["bloom"])
        q["co"] = "CO" + str((hash(q["bloom"]) % 5) + 1)

    return qs


def parse_syllabus_topics(syllabus: str) -> list[str]:
    if not syllabus:
        return []
    topics = []
    for part in re.split(r"[\n,;]", syllabus):
        cleaned = re.sub(r"\b(unit\s*\d+[:\-]?\b|topic\s*[:\-]?\b|chapter\s*[:\-]?\b)", "", part, flags=re.I)
        topic = cleaned.strip().strip(".-")
        if topic:
            topics.append(topic)
    return topics


def compute_syllabus_coverage(question_text: str, syllabus: str):
    topics = parse_syllabus_topics(syllabus)
    if not topics:
        return {"score": 100, "covered": [], "missing": [], "topics": []}

    combined = " ".join(question["text"] for question in parse_questions(question_text))
    lowered = combined.lower()
    covered = []
    missing = []
    for topic in topics:
        label = topic.lower().strip()
        if not label:
            continue
        if label in lowered or all(word in lowered for word in label.split()):
            covered.append(topic)
        else:
            missing.append(topic)

    score = round((len(covered) / len(topics)) * 100) if topics else 100
    return {"score": score, "covered": covered, "missing": missing, "topics": topics}


def validate_question_paper(text: str, questions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    seen = {}
    for q in questions:
        key = re.sub(r"[^a-z0-9]+", " ", q["text"].lower()).strip()
        if key:
            seen.setdefault(key, []).append(q["number"])
    for text_key, numbers in seen.items():
        if len(numbers) > 1:
            issues.append({
                "status": "warning",
                "title": "Duplicate question",
                "detail": f"Questions {numbers} appear to repeat the same wording.",
            })

    expected_numbers = [str(i) for i in range(1, len(questions) + 1)]
    actual_numbers = [str(q["number"]) for q in questions]
    if actual_numbers != expected_numbers:
        issues.append({
            "status": "warning",
            "title": "Question numbering",
            "detail": f"Questions are numbered as {actual_numbers}, expected {expected_numbers}.",
        })

    total_match = re.search(r"total\s+marks\s*[:\-]?\s*(\d+)", text, flags=re.I)
    if total_match:
        total_marks = int(total_match.group(1))
        extracted_marks = sum(int(q["marks"]) for q in questions)
        if extracted_marks != total_marks:
            issues.append({
                "status": "warning",
                "title": "Mark total mismatch",
                "detail": f"Question marks total {extracted_marks}, but total marks is {total_marks}.",
            })

    if not questions:
        issues.append({
            "status": "error",
            "title": "No questions found",
            "detail": "No valid question data could be extracted.",
        })

    return issues


def compute_co_balance(co_distribution: dict[str, int]) -> dict[str, Any]:
    levels = [f"CO{i}" for i in range(1, 6)]
    counts = {level: int(co_distribution.get(level, 0)) for level in levels}
    total = sum(counts.values())
    if total == 0:
        return {"score": 100, "counts": counts, "target": {level: 20 for level in levels}, "notes": "No CO data available."}

    ratios = {level: counts[level] / total for level in levels}
    deviations = [abs(ratios[level] - 0.2) for level in levels]
    balance_score = max(0, round((1 - (sum(deviations) / len(levels))) * 100))
    return {
        "score": balance_score,
        "counts": counts,
        "target": {level: 20 for level in levels},
        "notes": "CO distribution is balanced when each outcome is near 20% of total marks/questions.",
    }


def compute_rubric(coverage: dict[str, Any], validation: list[dict[str, Any]], co_balance: dict[str, Any], bloom_distribution: dict[str, int]) -> dict[str, Any]:
    bloom_levels = ["Remember", "Understand", "Apply", "Analyze", "Evaluate", "Create"]
    bloom_total = sum(bloom_distribution.values())
    if bloom_total:
        bloom_balance = min(100, round((max(bloom_distribution.values()) / bloom_total) * 100))
    else:
        bloom_balance = 100

    score = round(
        (coverage["score"] * 0.45)
        + ((100 - (len(validation) * 12)) * 0.35)
        + (co_balance["score"] * 0.20)
    )
    score = max(0, min(100, score))

    return {
        "score": score,
        "targets": {
            "coverage": 80,
            "validation": 0,
            "co_balance": 80,
            "bloom_mix": 60,
        },
        "checks": {
            "coverage": coverage["score"],
            "validation_issues": len(validation),
            "co_balance": co_balance["score"],
            "bloom_mix": bloom_balance,
        },
        "bloom_levels": {level: bloom_distribution.get(level, 0) for level in bloom_levels},
    }


def compute_blueprint(difficulty_distribution: dict[str, int]) -> dict[str, Any]:
    actual = {"Easy": int(difficulty_distribution.get("Easy", 0)), "Medium": int(difficulty_distribution.get("Medium", 0)), "Hard": int(difficulty_distribution.get("Hard", 0))}
    total = sum(actual.values())
    targets = {"Easy": 30, "Medium": 50, "Hard": 20}
    if total == 0:
        return {"targets": targets, "actual": actual, "score": 100, "status": "No difficulty data"}

    ratios = {level: (actual[level] / total) * 100 for level in ["Easy", "Medium", "Hard"]}
    deviation = sum(abs(ratios[level] - targets[level]) for level in ["Easy", "Medium", "Hard"]) / 3
    score = max(0, min(100, round(100 - deviation)))
    status = "Pass" if score >= 80 else "Needs review"
    return {"targets": targets, "actual": actual, "score": score, "status": status}


def compute_policy(blueprint: dict[str, Any], coverage: dict[str, Any], validation: list[dict[str, Any]], co_balance: dict[str, Any]) -> dict[str, Any]:
    checks = {
        "blueprint": blueprint["score"],
        "coverage": coverage["score"],
        "validation": max(0, 100 - (len(validation) * 12)),
        "co_balance": co_balance["score"],
    }
    score = round(sum(checks.values()) / len(checks))
    return {
        "score": max(0, min(100, score)),
        "status": "Pass" if score >= 75 else "Needs review",
        "checks": checks,
    }


def compare_audits(audit_a: dict[str, Any], audit_b: dict[str, Any]) -> dict[str, Any]:
    quality_delta = int(audit_b.get("quality_score", 0) - audit_a.get("quality_score", 0))
    coverage_delta = int(audit_b.get("coverage", {}).get("score", 0) - audit_a.get("coverage", {}).get("score", 0))
    rubric_delta = int(audit_b.get("rubric", {}).get("score", 0) - audit_a.get("rubric", {}).get("score", 0))
    policy_delta = int(audit_b.get("policy", {}).get("score", 0) - audit_a.get("policy", {}).get("score", 0))

    return {
        "left_id": audit_a.get("id"),
        "right_id": audit_b.get("id"),
        "delta": {
            "quality_score": quality_delta,
            "coverage": coverage_delta,
            "rubric": rubric_delta,
            "policy": policy_delta,
        },
        "side_by_side": {
            "quality": {
                "left": audit_a.get("quality_score", 0),
                "right": audit_b.get("quality_score", 0),
            },
            "coverage": {
                "left": audit_a.get("coverage", {}).get("score", 0),
                "right": audit_b.get("coverage", {}).get("score", 0),
            },
            "rubric": {
                "left": audit_a.get("rubric", {}).get("score", 0),
                "right": audit_b.get("rubric", {}).get("score", 0),
            },
            "policy": {
                "left": audit_a.get("policy", {}).get("score", 0),
                "right": audit_b.get("policy", {}).get("score", 0),
            },
        },
        "summary": {
            "left": audit_a.get("filename"),
            "right": audit_b.get("filename"),
            "winner": "right" if quality_delta >= 0 else "left",
        },
        "recommendation": "The revised paper is stronger overall; keep the blueprint and coverage patterns while improving weaker sections with more higher-order questions." if quality_delta >= 0 else "The newer paper is weaker; revisit coverage, blueprint balance, and deeper Bloom-level questions.",
    }


def generate_comparison_pdf(audit_a: dict[str, Any], audit_b: dict[str, Any]) -> bytes:
    if not FPDF:
        raise RuntimeError("fpdf2 not installed")

    comparison = compare_audits(audit_a, audit_b)
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, "Audit Comparison Report", ln=True)
    pdf.set_font("Arial", "", 12)
    pdf.cell(0, 8, f"Left: {audit_a.get('filename', 'left')}", ln=True)
    pdf.cell(0, 8, f"Right: {audit_b.get('filename', 'right')}", ln=True)
    pdf.ln(4)
    for key, value in comparison["delta"].items():
        pdf.cell(0, 8, f"{key}: {value}", ln=True)
    pdf.ln(4)
    pdf.multi_cell(0, 8, comparison["recommendation"])
    return pdf.output(dest="S")


def generate_recommendations(coverage: dict[str, Any], validation: list[dict[str, Any]], co_balance: dict[str, Any], bloom_distribution: dict[str, int], blueprint: dict[str, Any] | None = None) -> list[str]:
    recs: list[str] = []

    if coverage.get("missing"):
        missing = ", ".join(coverage["missing"][:3])
        recs.append(f"Coverage: add or revise questions on {missing} to strengthen syllabus alignment.")

    if validation:
        recs.append("Validation: review duplicate wording, numbering sequence, and total marks before finalizing the paper.")

    if co_balance.get("score", 100) < 80:
        recs.append("CO balance: redistribute marks across the CO outcomes so the paper reflects the full curriculum map more evenly.")

    higher_order = bloom_distribution.get("Analyze", 0) + bloom_distribution.get("Evaluate", 0) + bloom_distribution.get("Create", 0)
    if higher_order == 0:
        recs.append("Bloom mix: include more analyze, evaluate, and create prompts to lift the cognitive difficulty of the paper.")

    if blueprint and blueprint.get("score", 100) < 80:
        recs.append("Blueprint: rebalance the paper toward a 30% easy, 50% medium, 20% hard mix for a healthier exam design.")

    if not recs:
        recs.append("Assessment quality is strong. Keep the current mix of Bloom levels and preserve balanced CO coverage in future papers.")

    return recs


def _ensure_data_dir() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)


def audit_question_paper(text: str, filename: str, syllabus: str = ""):
    questions = parse_questions(text)
    if not questions:
        raise ValueError("No questions detected")

    coverage = compute_syllabus_coverage(text, syllabus)
    validation = validate_question_paper(text, questions)
    summary = {
        "question_count": len(questions),
        "extracted_marks": sum(int(q["marks"]) for q in questions),
        "coverage_score": coverage["score"],
        "validation_count": len(validation),
    }
    bloom_distribution = dict(Counter(q["bloom"] for q in questions))
    difficulty_distribution = dict(Counter(q["difficulty"] for q in questions))
    co_distribution = dict(Counter(q["co"] for q in questions))
    co_balance = compute_co_balance(co_distribution)
    blueprint = compute_blueprint(difficulty_distribution)
    policy = compute_policy(blueprint, coverage, validation, co_balance)
    rubric = compute_rubric(coverage, validation, co_balance, bloom_distribution)
    recommendations = generate_recommendations(coverage, validation, co_balance, bloom_distribution, blueprint)

    coverage_penalty = max(0, 100 - coverage["score"])
    quality_score = round(min(100, max(0, rubric["score"] - (coverage_penalty * 0.1))))

    findings = [{
        "status": "good",
        "title": "Success",
        "detail": f"Found {len(questions)} questions and {summary['extracted_marks']} marks.",
    }]
    if coverage["missing"]:
        findings.append({
            "status": "warning",
            "title": "Coverage gap",
            "detail": "Missing syllabus topics: " + ", ".join(coverage["missing"][:3]) + ("..." if len(coverage["missing"]) > 3 else "") + ".",
        })
    for issue in validation:
        findings.append({
            "status": issue["status"],
            "title": issue["title"],
            "detail": issue["detail"],
        })
    if syllabus:
        findings.append({
            "status": "info",
            "title": "Syllabus supplied",
            "detail": "The syllabus field was used for topic highlighting and coverage checks.",
        })
    findings.append({
        "status": "info",
        "title": "Recommended next steps",
        "detail": recommendations[0] if recommendations else "No additional recommendations.",
    })

    return {
        "id": uuid.uuid4().hex[:8],
        "filename": filename,
        "created_at": _iso_timestamp(),
        "quality_score": quality_score,
        "questions": questions,
        "summary": summary,
        "coverage": coverage,
        "validation": validation,
        "co_balance": co_balance,
        "rubric": rubric,
        "blueprint": blueprint,
        "policy": policy,
        "recommendations": recommendations,
        "bloom_distribution": bloom_distribution,
        "difficulty_distribution": difficulty_distribution,
        "co_distribution": co_distribution,
        "findings": findings,
    }


def save_audit(data: dict[str, Any]) -> None:
    _ensure_data_dir()
    conn = sqlite3.connect(str(AUDIT_DB))
    try:
        conn.execute(
            "CREATE TABLE IF NOT EXISTS audits (id TEXT PRIMARY KEY, filename TEXT, payload TEXT, created_at TEXT)"
        )
        columns = {row[1] for row in conn.execute("PRAGMA table_info(audits)").fetchall()}
        if "filename" not in columns:
            conn.execute("ALTER TABLE audits ADD COLUMN filename TEXT")
        if "payload" not in columns:
            conn.execute("ALTER TABLE audits ADD COLUMN payload TEXT")
        if "created_at" not in columns:
            conn.execute("ALTER TABLE audits ADD COLUMN created_at TEXT")
        conn.execute(
            "INSERT INTO audits (id, filename, payload, created_at) VALUES (?, ?, ?, ?)",
            (data["id"], data.get("filename"), json.dumps(data), data["created_at"]),
        )
        conn.commit()
    finally:
        conn.close()

    if AUDIT_STORE.exists():
        try:
            existing = json.loads(AUDIT_STORE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            existing = {}
    else:
        existing = {}

    if not isinstance(existing, dict):
        existing = {}
    existing[data["id"]] = data
    AUDIT_STORE.write_text(json.dumps(existing, indent=2), encoding="utf-8")


def get_audit(audit_id: str) -> dict[str, Any]:
    try:
        conn = sqlite3.connect(str(AUDIT_DB))
        try:
            row = conn.execute("SELECT payload FROM audits WHERE id = ?", (audit_id,)).fetchone()
        finally:
            conn.close()
    except sqlite3.OperationalError:
        row = None

    if row:
        return json.loads(row[0])

    if AUDIT_STORE.exists():
        try:
            data = json.loads(AUDIT_STORE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            data = {}
        if isinstance(data, dict) and audit_id in data:
            return data[audit_id]

    raise KeyError(f"No audit found for id {audit_id}")


def list_audits() -> list[dict[str, Any]]:
    audits: list[dict[str, Any]] = []
    try:
        conn = sqlite3.connect(str(AUDIT_DB))
        try:
            rows = conn.execute(
                "SELECT id, payload FROM audits ORDER BY created_at DESC"
            ).fetchall()
        finally:
            conn.close()
        audits = [json.loads(row[1]) for row in rows]
    except sqlite3.OperationalError:
        audits = []

    if not audits and AUDIT_STORE.exists():
        try:
            legacy = json.loads(AUDIT_STORE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            legacy = {}
        if isinstance(legacy, dict):
            audits = list(legacy.values())

    return sorted(audits, key=lambda item: item.get("created_at", ""), reverse=True)


@app.post("/api/audit")
async def run_audit(file: UploadFile = File(...), syllabus: str = Form("")):
    try:
        content = await file.read()
        text = extract_text(file.filename, content)
        audit = audit_question_paper(text, file.filename, syllabus)
        save_audit(audit)
        return audit
    except Exception as exc:  # pragma: no cover - surfaced as HTTP 500
        raise HTTPException(500, str(exc)) from exc


@app.get("/api/reports")
def list_reports():
    try:
        return list_audits()
    except Exception:
        return []


@app.get("/api/reports/{audit_id}")
def get_report(audit_id: str):
    try:
        return get_audit(audit_id)
    except KeyError as exc:
        raise HTTPException(404, "Report not found") from exc


@app.get("/api/reports/compare/{left_id}/{right_id}")
def compare_reports(left_id: str, right_id: str):
    try:
        left = get_audit(left_id)
        right = get_audit(right_id)
        return compare_audits(left, right)
    except KeyError as exc:
        raise HTTPException(404, "One or both reports were not found") from exc


@app.get("/api/reports/compare/{left_id}/{right_id}/pdf")
def get_comparison_pdf(left_id: str, right_id: str):
    try:
        left = get_audit(left_id)
        right = get_audit(right_id)
    except KeyError as exc:
        raise HTTPException(404, "One or both reports were not found") from exc

    try:
        pdf_bytes = generate_comparison_pdf(left, right)
    except RuntimeError as exc:
        raise HTTPException(500, str(exc)) from exc
    return StreamingResponse(io.BytesIO(pdf_bytes), media_type="application/pdf")


@app.get("/api/reports/{audit_id}/pdf")
def get_pdf(audit_id: str):
    try:
        data = get_audit(audit_id)
    except KeyError as exc:
        raise HTTPException(404, "Report not found") from exc

    if not FPDF:
        raise HTTPException(500, "fpdf2 not installed")

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, f"Audit: {data['filename']}", ln=True)
    pdf.set_font("Arial", "", 12)
    for q in data["questions"]:
        pdf.multi_cell(0, 10, f"{q['number']}. {q['text']} [{q['bloom']}]")
    return StreamingResponse(io.BytesIO(pdf.output(dest="S")), media_type="application/pdf")


@app.get("/")
def index():
    return FileResponse("index.html")
